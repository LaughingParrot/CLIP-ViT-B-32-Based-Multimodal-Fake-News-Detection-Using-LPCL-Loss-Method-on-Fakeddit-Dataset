from __future__ import annotations

import argparse
import os
from pathlib import Path
from typing import Iterable, Iterator, TextIO

EXCLUDED_DIRS = frozenset({
    "__pycache__",
    "venv",
    "env",
    ".venv",
    "node_modules",
    ".git",
    ".idea",
    ".vscode",
    "build",
    "dist",
    "site-packages",
    ".pytest_cache",
})

EXCLUDED_EXTENSIONS = frozenset({
    ".pyc",
    ".pyo",
    ".pyd",
    ".so",
    ".dll",
    ".exe",
    ".bin",
    ".obj",
    ".class",
    ".png",
    ".jpg",
    ".jpeg",
    ".gif",
    ".ico",
    ".pdf",
    ".svg",
    ".mp3",
    ".mp4",
    ".zip",
    ".tar",
    ".gz",
    ".rar",
    ".7z",
    ".sqlite3",
    ".db",
})

PYTHON_EXTS = frozenset({".py", ".pyw"})
SPREADSHEET_EXTS = frozenset({".csv", ".tsv", ".xls", ".xlsx"})
TEXT_EXTS = frozenset({
    ".txt",
    ".md",
    ".json",
    ".yaml",
    ".yml",
    ".ini",
    ".cfg",
    ".doc",
    ".docx",
    ".rtf",
    ".log",
})

TREE_BRANCH = "|-- "
TREE_LAST = "`-- "
TREE_PIPE = "|   "
TREE_GAP = "    "
HEADER_LINE = "=" * 60
SECTION_LINE = "-" * 60
SOURCE_LINE = "*" * 80


def _is_ignored_name(name: str, *, is_dir: bool, suffix: str = "") -> bool:
    normalized_name = name.casefold()

    if normalized_name.startswith("."):
        return True

    if normalized_name in EXCLUDED_DIRS:
        return True

    if not is_dir and suffix.lower() in EXCLUDED_EXTENSIONS:
        return True

    return False


def is_ignored(path: Path) -> bool:
    try:
        is_dir = path.is_dir()
    except OSError:
        # If metadata cannot be read, safest behavior is to skip.
        return True

    return _is_ignored_name(path.name, is_dir=is_dir, suffix=path.suffix)


def _iter_tree_lines(dir_path: Path, prefix: str = "") -> Iterator[str]:
    entries: list[tuple[Path, bool]] = []

    try:
        for path in dir_path.iterdir():
            try:
                is_dir = path.is_dir()
            except OSError:
                continue

            if _is_ignored_name(path.name, is_dir=is_dir, suffix=path.suffix):
                continue

            entries.append((path, is_dir))
    except PermissionError:
        yield f"{prefix}{TREE_LAST}[Access Denied]\n"
        return
    except OSError as exc:
        yield f"{prefix}{TREE_LAST}[Error: {exc}]\n"
        return

    entries.sort(key=lambda item: (not item[1], item[0].name.casefold()))

    last_index = len(entries) - 1
    for index, (path, is_dir) in enumerate(entries):
        is_last = index == last_index
        marker = TREE_LAST if is_last else TREE_BRANCH

        yield f"{prefix}{marker}{path.name}\n"

        if is_dir:
            extension = TREE_GAP if is_last else TREE_PIPE
            yield from _iter_tree_lines(path, prefix=prefix + extension)


def generate_tree(dir_path: Path, prefix: str = "") -> str:
    return "".join(_iter_tree_lines(dir_path, prefix))


def _iter_content_file_paths(target_path: Path, allowed_exts: frozenset[str]) -> Iterator[Path]:
    for root, dirs, files in os.walk(target_path, topdown=True):
        dirs[:] = sorted(
            (
                directory
                for directory in dirs
                if not _is_ignored_name(directory, is_dir=True)
            ),
            key=str.casefold,
        )

        for filename in sorted(files, key=str.casefold):
            suffix = Path(filename).suffix.lower()

            if suffix not in allowed_exts:
                continue

            if _is_ignored_name(filename, is_dir=False, suffix=suffix):
                continue

            yield Path(root) / filename


def _normalize_path(path: Path) -> str:
    return os.path.normcase(os.path.abspath(path))


def iter_extracted_contents(
    target_path: Path,
    allowed_exts: frozenset[str],
    *,
    base_path: Path,
    skip_paths: frozenset[str],
) -> Iterator[str]:
    for file_path in _iter_content_file_paths(target_path, allowed_exts):
        if _normalize_path(file_path) in skip_paths:
            continue
        yield _read_file(file_path, base_path)


def extract_contents(
    target_path: Path,
    allowed_exts: frozenset[str],
    *,
    base_path: Path,
    skip_paths: frozenset[str],
) -> str:
    return "".join(
        iter_extracted_contents(
            target_path,
            allowed_exts,
            base_path=base_path,
            skip_paths=skip_paths,
        )
    )


def _relative_display_path(file_path: Path, base_path: Path) -> str:
    for anchor in (base_path, Path.cwd()):
        try:
            return str(file_path.relative_to(anchor))
        except ValueError:
            continue

    return file_path.name


def _read_file(file_path: Path, base_path: Path) -> str:
    ext = file_path.suffix.lower()
    rel_path = _relative_display_path(file_path, base_path)
    header = f"\n{HEADER_LINE}\nFile: {rel_path}\n{HEADER_LINE}\n\n"

    try:
        if ext == ".xlsx":
            text = _read_xlsx(file_path)
        elif ext == ".docx":
            text = _read_docx(file_path)
        elif ext in {".doc", ".xls"}:
            text = f"[Legacy binary format {ext} skipped.]"
        else:
            text = file_path.read_text(encoding="utf-8", errors="replace")

        return header + text + "\n"

    except PermissionError:
        return header + "[Error: Permission Denied. File is locked.]\n"
    except OSError as exc:
        return header + f"[Error reading file: {exc}]\n"


def _read_xlsx(file_path: Path) -> str:
    try:
        import openpyxl
    except ImportError:
        return "[Requires 'openpyxl' library.]"

    try:
        workbook = openpyxl.load_workbook(file_path, data_only=True, read_only=True)
    except Exception as exc:
        return f"[Failed to parse Excel file: {exc}]"

    lines: list[str] = []

    try:
        for sheet in workbook.worksheets:
            lines.append(f"\n--- Sheet: {sheet.title} ---")

            for row in sheet.iter_rows(values_only=True):
                if not row:
                    continue

                values = ["" if cell is None else str(cell) for cell in row]
                if any(value.strip() for value in values):
                    lines.append("\t".join(values))
    except Exception as exc:
        return f"[Failed to parse Excel file: {exc}]"
    finally:
        workbook.close()

    return "\n".join(lines)


def _read_docx(file_path: Path) -> str:
    try:
        import docx
    except ImportError:
        return "[Requires 'python-docx' library.]"

    try:
        doc = docx.Document(str(file_path))
        return "\n".join(paragraph.text for paragraph in doc.paragraphs if paragraph.text)
    except Exception as exc:
        return f"[Failed to parse Word file: {exc}]"


def _build_allowed_extensions(*, include_spreadsheets: bool, include_text: bool) -> frozenset[str]:
    allowed = set(PYTHON_EXTS)

    if include_spreadsheets:
        allowed.update(SPREADSHEET_EXTS)

    if include_text:
        allowed.update(TEXT_EXTS)

    return frozenset(allowed)


def _write_header(output_handle: TextIO, text: str) -> None:
    output_handle.write(text)


def _write_content_blocks(output_handle: TextIO, blocks: Iterable[str]) -> None:
    for block in blocks:
        output_handle.write(block)


def main() -> None:
    parser = argparse.ArgumentParser(description="Extract codebase and document content cleanly.")
    parser.add_argument(
        "paths",
        nargs="*",
        default=["."],
        help="Target files or directories. Defaults to current directory.",
    )
    parser.add_argument("-o", "--output", default="code_dump.txt", help="Output file path.")

    group = parser.add_mutually_exclusive_group()
    group.add_argument("--tree-only", action="store_true", help="Output ONLY the directory tree.")
    group.add_argument("--content-only", action="store_true", help="Output ONLY the file contents.")

    parser.add_argument(
        "-s",
        "--spreadsheets",
        action="store_true",
        help="Include spreadsheets (.csv, .xlsx, etc.)",
    )
    parser.add_argument(
        "-t",
        "--text",
        action="store_true",
        help="Include text/docs (.txt, .md, .docx, etc.)",
    )

    args = parser.parse_args()

    allowed_exts = _build_allowed_extensions(
        include_spreadsheets=args.spreadsheets,
        include_text=args.text,
    )

    output_path = Path(args.output).resolve()
    skip_paths = frozenset({_normalize_path(output_path)})

    try:
        with output_path.open("w", encoding="utf-8", newline="\n") as output_handle:
            for path_str in args.paths:
                target_path = Path(path_str).resolve()

                if not target_path.exists():
                    print(f"Warning: '{target_path}' does not exist. Skipping.")
                    continue

                _write_header(output_handle, f"Source Path: {target_path}\n\n")

                if target_path.is_dir():
                    if not args.content_only:
                        _write_header(output_handle, "### DIRECTORY TREE ###\n")
                        _write_header(output_handle, f"{target_path.name}/\n")
                        _write_header(output_handle, generate_tree(target_path))
                        _write_header(output_handle, f"\n{SECTION_LINE}\n\n")

                    if not args.tree_only:
                        _write_header(output_handle, "### FILE CONTENTS ###\n")
                        _write_content_blocks(
                            output_handle,
                            iter_extracted_contents(
                                target_path,
                                allowed_exts,
                                base_path=target_path,
                                skip_paths=skip_paths,
                            ),
                        )

                elif target_path.is_file() and not args.tree_only:
                    if _normalize_path(target_path) in skip_paths:
                        print(f"Warning: '{target_path}' is the output file. Skipping self-read.")
                    else:
                        _write_header(output_handle, "### FILE CONTENTS ###\n")
                        _write_header(output_handle, _read_file(target_path, Path.cwd()))

                _write_header(output_handle, f"\n{SOURCE_LINE}\n\n")

        print(f"Extraction complete. Saved to {output_path}")
    except OSError as exc:
        print(f"Error writing to output file: {exc}")


if __name__ == "__main__":
    main()
